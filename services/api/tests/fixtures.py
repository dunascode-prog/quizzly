"""Builds small real PDF and PPTX files for tests, so extraction runs on genuine documents."""

import io

from pptx import Presentation
from pptx.util import Inches


def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]]) -> bytes:
    """One PDF page per list item; each string is a line of Helvetica text."""
    objects: list[bytes] = []
    n = len(pages)
    # 1: catalog, 2: page tree, 3: font, then (page, content) pairs.
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, lines in enumerate(pages):
        ops = ["BT", "/F1 12 Tf", "14 TL", "72 760 Td"]
        ops += [f"({_pdf_escape(line)}) Tj T*" for line in lines]
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for num, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{num} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_pptx(slides: list[tuple[str, list[str]]]) -> bytes:
    """(title, bullet lines) per slide, using the 'Title and Content' layout."""
    prs = Presentation()
    for title, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0] if bullets else ""
        for line in bullets[1:]:
            body.add_paragraph().text = line
        if title == "Table":
            shape = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1))
            shape.table.cell(0, 0).text = "Push"
            shape.table.cell(0, 1).text = "O(1)"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
