"""Per-page text from PDF, PPTX and photo sets (ARCHITECTURE §5.1). OCR itself is in ocr.py."""

import io
import re
import zipfile
from dataclasses import dataclass

import pypdfium2 as pdfium
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.exc import PackageNotFoundError

PDF_MIME = "application/pdf"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
# 1–30 photos uploaded as one document; each photo is a page (FR-DOC-01).
IMAGES_MIME = "application/x-quizzly-images"
# A PDF page with less text than this is treated as a scan and OCR'd.
SCANNED_PAGE_CHARS = 20


@dataclass
class Page:
    number: int
    text: str
    # A slide title or a likely heading; drives the outline when a deck has them (FR-TOP-02).
    heading: str | None = None
    method: str = "text"  # pages.extraction_method: "text" or "ocr"
    ocr_confidence: float | None = None
    low_confidence: bool = False


class ExtractionError(Exception):
    """A file we can't read. `reason` is a source_files.failure_reason value."""

    def __init__(self, reason: str, detail: str = ""):
        super().__init__(detail or reason)
        self.reason = reason


def sniff(data: bytes) -> str:
    """Decide the format from the bytes, not the file name (students rename files)."""
    if data[:5] == b"%PDF-":
        return PDF_MIME
    if data[:4] == b"PK\x03\x04":
        return PPTX_MIME
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        # Office's container for password-protected (and old .ppt) files.
        raise ExtractionError("encrypted", "OLE container: encrypted or legacy Office file")
    raise ExtractionError("corrupt", "unrecognised file format")


def check_image(data: bytes) -> None:
    """A photo must be a real JPEG or PNG; anything else fails the document as corrupt."""
    try:
        with Image.open(io.BytesIO(data)) as img:
            if img.format not in ("JPEG", "PNG", "MPO"):  # MPO: some phones' multi-picture JPEG
                raise ExtractionError("corrupt", f"unsupported image format {img.format}")
            img.verify()
    except (UnidentifiedImageError, OSError, SyntaxError) as exc:
        raise ExtractionError("corrupt", repr(exc)) from exc


def render_pdf_pages(data: bytes, numbers: list[int], scale: float = 2.5):
    """Yields (page number, PIL image) for the given 1-based pages, for OCR of scans."""
    pdf = _open_pdf(data)
    try:
        for n in numbers:
            page = pdf[n - 1]
            try:
                yield n, page.render(scale=scale).to_pil()
            finally:
                page.close()
    finally:
        pdf.close()


def page_count(data: bytes, mime: str) -> int:
    if mime == PDF_MIME:
        pdf = _open_pdf(data)
        try:
            return len(pdf)
        finally:
            pdf.close()
    return len(_open_pptx(data).slides)


def extract(data: bytes, mime: str) -> list[Page]:
    return _extract_pdf(data) if mime == PDF_MIME else _extract_pptx(data)


def _open_pdf(data: bytes) -> pdfium.PdfDocument:
    try:
        return pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        reason = "encrypted" if "password" in str(exc).lower() else "corrupt"
        raise ExtractionError(reason, str(exc)) from exc


def _extract_pdf(data: bytes) -> list[Page]:
    pdf = _open_pdf(data)
    try:
        pages = []
        for i in range(len(pdf)):
            page = pdf[i]
            textpage = page.get_textpage()
            try:
                text = textpage.get_text_bounded()
            finally:
                textpage.close()
                page.close()
            pages.append(Page(number=i + 1, text=text))
        return pages
    finally:
        pdf.close()


def _open_pptx(data: bytes):
    try:
        return Presentation(io.BytesIO(data))
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError, ValueError) as exc:
        raise ExtractionError("corrupt", repr(exc)) from exc


def _shape_texts(shapes) -> list[str]:
    out: list[str] = []
    for shape in shapes:
        if shape.shape_type == 6:  # group: recurse
            out.extend(_shape_texts(shape.shapes))
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            out.append(shape.text_frame.text)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                out.append(" | ".join(cell.text for cell in row.cells))
    return [t for t in out if t.strip()]


def _extract_pptx(data: bytes) -> list[Page]:
    prs = _open_pptx(data)
    pages = []
    for i, slide in enumerate(prs.slides):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        body = [t for t in _shape_texts(slide.shapes) if t.strip() != title]
        text = "\n".join([title, *body]) if title else "\n".join(body)
        pages.append(Page(number=i + 1, text=text, heading=_one_line(title) or None))
    return pages


def _one_line(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()[:120]
